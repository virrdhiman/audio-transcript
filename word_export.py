"""Export transcripts to Microsoft Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from transcriber import TranscriptResult, format_duration, format_timestamp


def export_to_word(
    result: TranscriptResult,
    output_path: str | Path,
    source_name: str,
) -> Path:
    """Write a Word document with call duration and word-by-word timestamps."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    title = doc.add_heading("Call / Audio Transcript", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run("Source file: ").bold = True
    meta.add_run(source_name)

    duration_p = doc.add_paragraph()
    duration_p.add_run("Call recording time: ").bold = True
    duration_p.add_run(
        f"{format_duration(result.duration_seconds)} "
        f"({format_timestamp(result.duration_seconds)})"
    )

    if result.language:
        lang_p = doc.add_paragraph()
        lang_p.add_run("Detected language: ").bold = True
        lang_p.add_run(result.language)

    doc.add_heading("Full Transcript", level=1)
    full = doc.add_paragraph(result.full_text or "(No speech detected)")
    for run in full.runs:
        run.font.size = Pt(11)

    doc.add_heading("Word-by-Word Transcript with Timestamps", level=1)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("#", "Timestamp", "Word", "End")
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for index, word in enumerate(result.words, start=1):
        row = table.add_row().cells
        row[0].text = str(index)
        row[1].text = format_timestamp(word.start)
        row[2].text = word.word
        row[3].text = format_timestamp(word.end)

    # Keep table readable in Word
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()
    timed = doc.add_paragraph()
    timed.add_run("Inline timed transcript:").bold = True

    inline = doc.add_paragraph()
    if not result.words:
        inline.add_run("(No words detected)")
    else:
        for word in result.words:
            stamp = inline.add_run(f"[{format_timestamp(word.start)}] ")
            stamp.bold = True
            stamp.font.size = Pt(9)
            token = inline.add_run(f"{word.word} ")
            token.font.size = Pt(11)

    section = doc.sections[0]
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    doc.save(str(out))
    return out
